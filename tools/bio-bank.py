# bio-bank.py — 从生物真题解析卷/答案卷 docx 批量抽取选择题库
# 用法：python tools/bio-bank.py <真题根目录> <输出.json>
#
# 两种常见排版都要吃：
#   A. 内联答案：题干 -> A./B./C./D. -> 【答案】X -> 【详解】...
#   B. 卷末答案块：题干 -> 选项（无答案） ... 「生物·答案」 -> 1-5 BDDCB 6-10 DACAC ...
import io, json, os, re, sys, glob

import docx

Q_START = re.compile(r"^(\d{1,2})\s*[.．、]\s*(.+)$")
OPT_LINE = re.compile(r"^\s*([A-D])\s*[.．、]\s*(.*)$")
# 北京卷有 "A 乙烯抑制主根生长" 这种不带点号的选项
OPT_PLAIN = re.compile(r"^([A-D])[\s\u3000]+(\S.*)$")
# 关键：选项字母前面不能是英文字母/数字，否则 DNA、ATP、C3 里的 A/C 会被误切
OPT_INLINE = re.compile(r"(?<![A-Za-z0-9])([A-D])\s*[.．、]\s*(.*?)(?=\s*(?<![A-Za-z0-9])[A-D]\s*[.．、]|$)")
ANS_INLINE = re.compile(r"【答案】\s*([A-D]+)")
YEAR = re.compile(r"(\d{4})年")
RANGE = re.compile(r"(\d{1,2})\s*[-–~]\s*(\d{1,2})\s*[.．、]?\s*([A-D]{2,})")
SINGLE = re.compile(r"(?<![\dA-Da-d])(\d{1,2})\s*[.．、]\s*([A-D]{1,4})(?![A-Za-z])")
PAGENO = re.compile(r"^第\s*\d+\s*页|共\s*\d+\s*页|^\s*\d+\s*/\s*\d+\s*$")


def area_of(name):
    m = re.search(r"（([^（）]+?)）", name)
    if not m:
        return ""
    a = m.group(1)
    a = re.sub(r"\((空白卷|解析卷|答案|原卷|精编.*?)\)", "", a).strip()
    mm = re.search(r"【(\d+)月】", name)
    if mm and a.startswith("浙江"):
        a = "浙江%s月" % mm.group(1)
    return a.rstrip("卷") or a


def split_options(text):
    """切出 A/B/C/D 标记。

    北京卷有两种写法：段落开头 "A 乙烯抑制主根生长"（无点号）与 "A. 自养"
    （有点号），还有一行里混排的。做法是先把段首的无点号标记补成 "A. "，
    再统一按点号切——这样题干里 "A 与 T 配对" 这类不算段首，不会被误切。

    关键补丁：标记必须按 A→B→C→D 顺序出现才收。否则 "元素是 C、H、O" 里的
    C 会被当成选项 C，把一道题切得七零八落（2024 江苏第 1 题就这么坏的）。
    """
    m = re.match(r"^([A-D])[\s\u3000]+", text)
    if m:
        text = text[0] + ". " + text[m.end():]
    m = re.match(r"^([A-D])[\s\u3000]+", text)
    if m:
        text = text[0] + ". " + text[m.end():]
    # 迭代：遇到顺序不对的「假标记」（如元素符号 C、），去掉它的标点当普通文字，
    # 再重切，直到切出的标记严格递增或够 4 个
    for _ in range(6):
        out, reject = [], None
        for mm in OPT_INLINE.finditer(text):
            if len(out) >= 4:
                break
            L = mm.group(1)
            if not out:
                if text[:mm.start()].strip():
                    reject = mm
                    break
                out.append((L, mm.group(2).strip()))
            elif "ABCD"["ABCD".index(out[0][0]) + len(out)] == L:
                out.append((L, mm.group(2).strip()))
            else:
                reject = mm
                break
        if reject is None:
            return out
        text = text[:reject.start()] + reject.group(1) + " " + text[reject.start(2):]
    return out


def parse_answers_block(ps, start):
    ans = {}
    for t in ps[start:]:
        for m in RANGE.finditer(t):
            lo, hi = int(m.group(1)), int(m.group(2))
            letters = m.group(3)
            for k in range(lo, hi + 1):
                idx = k - lo
                if idx < len(letters):
                    ans[k] = letters[idx]
        for m in SINGLE.finditer(t):
            ans.setdefault(int(m.group(1)), m.group(2))
    return ans


def parse_docx(path):
    name = os.path.basename(path)
    ym = YEAR.search(name)
    year = ym.group(1) if ym else ""
    area = area_of(name)
    ps = [p.text.replace("\u200b", "").replace("\xa0", " ").strip()
          for p in docx.Document(path).paragraphs]
    ps = [p for p in ps if p and not PAGENO.search(p)]

    qs, cur = [], None
    for i, t in enumerate(ps):
        m = Q_START.match(t)
        if m:
            no = int(m.group(1))
            if cur is None or no == cur["no"] + 1:
                cur = {"no": no, "stem": m.group(2).strip(), "opts": [],
                       "ans": "", "explain": "", "para": i}
                qs.append(cur)
                continue
        if cur is None:
            continue
        if ANS_INLINE.search(t):
            cur["ans"] = ANS_INLINE.search(t).group(1)
            continue
        if t.startswith("【"):
            cur["explain"] += t + " "
            continue
        if cur["ans"]:
            cur["explain"] += t + " "
            continue
        if re.match(r"^(第[一二三]部分|本部分共|一、|二、|三、|注意事项)", t) or \
                re.search(r"(2B铅笔|答题卡|考试结束后|本试卷共|考试时长)", t):
            continue
        opts = split_options(t)
        om = OPT_LINE.match(t) or OPT_PLAIN.match(t)
        if om and not opts:
            opts = [(om.group(1), om.group(2).strip())]
        if opts:
            cur["opts"].extend(opts)
        elif cur["opts"]:
            last_k, last_v = cur["opts"][-1]
            cur["opts"][-1] = (last_k, last_v + " " + t)
        else:
            cur["stem"] += " " + t

    if any(not q["ans"] for q in qs):
        start = 0
        for i, t in enumerate(ps):
            if re.search(r"(·答案|答案】|参考答案)", t) and i > len(ps) * 0.4:
                start = i
                break
        if start:
            block = parse_answers_block(ps, start)
            for q in qs:
                if not q["ans"] and q["no"] in block:
                    q["ans"] = block[q["no"]]

    rows = []
    for q in qs:
        if len(q["opts"]) < 4 or not q["ans"]:
            continue
        rows.append({
            "year": year, "area": area, "no": q["no"],
            "stem": re.sub(r"\s+", " ", q["stem"]).strip(),
            "options": [o[1] for o in q["opts"][:4]],
            "answer": q["ans"],
            "explain": re.sub(r"\s+", " ", q["explain"]).strip()[:1500],
            "material": "",
            "doc": path, "para": q["para"],
        })
    return year, area, rows, len(qs)


def parse_pair(blank_path, ans_path):
    """空白卷取题干选项，答案卷取「第 N 题 / 答案：X / 解析：…」。"""
    def paras(p):
        out = [x.text.replace("\u200b", "").replace("\xa0", " ").strip()
               for x in docx.Document(p).paragraphs]
        return [x for x in out if x and not PAGENO.search(x)]

    ps = paras(blank_path)
    name = os.path.basename(blank_path)
    ym = YEAR.search(name)
    year = ym.group(1) if ym else ""
    area = area_of(name)

    qs, cur = [], None
    for i, t in enumerate(ps):
        m = Q_START.match(t)
        if m:
            no = int(m.group(1))
            if cur is None or no == cur["no"] + 1:
                cur = {"no": no, "stem": m.group(2).strip(), "opts": [], "para": i}
                qs.append(cur)
                continue
        if cur is None:
            continue
        if re.match(r"^(第[一二三]部分|本部分共|一、|二、|三、|注意事项)", t) or \
                re.search(r"(2B铅笔|答题卡|考试结束后|本试卷共|考试时长)", t):
            continue
        opts = split_options(t)
        om = OPT_LINE.match(t) or OPT_PLAIN.match(t)
        if om and not opts:
            opts = [(om.group(1), om.group(2).strip())]
        if opts:
            cur["opts"].extend(opts)
        elif cur["opts"]:
            last_k, last_v = cur["opts"][-1]
            cur["opts"][-1] = (last_k, last_v + " " + t)
        else:
            cur["stem"] += " " + t

    # 答案卷
    aps = paras(ans_path)
    ans, expl = {}, {}
    curno = None
    for t in aps:
        mm = re.match(r"^第\s*(\d{1,2})\s*题", t)
        if mm:
            curno = int(mm.group(1))
            rest = t[mm.end():].strip()
            if rest:
                am = re.search(r"答案[：:]\s*([A-D]+)", rest)
                if am:
                    ans[curno] = am.group(1)
            continue
        if curno is None:
            continue
        am = re.search(r"答案[：:]\s*([A-D]+)", t)
        if am and curno not in ans:
            ans[curno] = am.group(1)
        # 答案与解析常常挤在同一段，别用 continue 跳过（否则解析全空）
        dm = re.search(r"解析[：:]\s*(.*)", t)
        if dm:
            expl[curno] = expl.get(curno, "") + dm.group(1) + " "
        elif t.startswith("解析"):
            continue
        elif not am:
            expl[curno] = expl.get(curno, "") + t + " "

    rows = []
    for q in qs:
        if len(q["opts"]) < 4 or q["no"] not in ans:
            continue
        rows.append({
            "year": year, "area": area, "no": q["no"],
            "stem": re.sub(r"\s+", " ", q["stem"]).strip(),
            "options": [o[1] for o in q["opts"][:4]],
            "answer": ans[q["no"]],
            "explain": re.sub(r"\s+", " ", expl.get(q["no"], "")).strip()[:1500],
            "material": "",
            "doc": blank_path, "para": q["para"],
        })
    return year, area, rows, len(qs)


def main():
    root, out_path = sys.argv[1], sys.argv[2]
    only = sys.argv[3] if len(sys.argv) > 3 else "生物"
    if only not in ("", "-"):
        flt = lambda x: only in os.path.basename(x)
    else:
        flt = lambda x: True

    files = sorted(glob.glob(os.path.join(root, "**", "*解析卷*.docx"), recursive=True))
    files += sorted(glob.glob(os.path.join(root, "**", "*答案*.docx"), recursive=True))
    files = sorted(set(x for x in files if flt(x)))
    print("解析卷/答案卷 docx:", len(files))
    rows, stats, seen = [], [], set()
    for f in files:
        try:
            year, area, rs, nq = parse_docx(f)
        except Exception as e:
            stats.append((os.path.basename(f), -1, 0, repr(e)[:50]))
            continue
        stats.append((os.path.basename(f), nq, len(rs), ""))
        for r in rs:
            key = (r["year"], r["area"], r["no"], r["stem"][:24])
            if key in seen:
                continue
            seen.add(key)
            rows.append(r)

    # 空白卷 + 答案卷 配对（北京 2026 这类只有空白卷+答案卷的年份）
    pairs = 0
    for blank in sorted(glob.glob(os.path.join(root, "**", "*空白卷*.docx"), recursive=True)):
        if not flt(blank):
            continue
        d, n = os.path.dirname(blank), os.path.basename(blank)
        cand = [x for x in glob.glob(os.path.join(d, "*.docx"))
                if x != blank and ("答案" in os.path.basename(x) or "解析" in os.path.basename(x))
                and os.path.basename(x)[:4] == n[:4]]
        if not cand:
            continue
        pick = cand[0]
        for c in cand:
            if "答案" in os.path.basename(c):
                pick = c
                break
        try:
            year, area, rs, nq = parse_pair(blank, pick)
        except Exception as e:
            stats.append((n, -1, 0, "pair:" + repr(e)[:40]))
            continue
        stats.append((n + " [配对]", nq, len(rs), ""))
        pairs += 1
        for r in rs:
            key = (r["year"], r["area"], r["no"], r["stem"][:24])
            if key in seen:
                continue
            seen.add(key)
            rows.append(r)
    print("空白+答案配对:", pairs)

    io.open(out_path, "w", encoding="utf-8").write(json.dumps(rows, ensure_ascii=False, indent=1))
    print("入库选择题:", len(rows))
    bad = [s for s in stats if s[1] < 0 or s[2] < max(1, s[1] * 0.6)]
    print("成题率偏低文件:", len(bad))
    for s in bad[:12]:
        print("   %-54s 题号%d 成题%d %s" % (s[0][:54], s[1], s[2], s[3]))
    by_year = {}
    for r in rows:
        by_year[r["year"]] = by_year.get(r["year"], 0) + 1
    print("按年份:", dict(sorted(by_year.items())))
    import collections
    bj = collections.Counter(r["year"] for r in rows if r["area"] == "北京")
    print("北京各年:", dict(sorted(bj.items())))


if __name__ == "__main__":
    main()
